from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from dashboard.i18n import get_text
from reports.base import ReportGenerator
from reports.context import ReportContext

ENCABEZADO_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
ENCABEZADO_FONT = Font(bold=True, color="FFFFFF")
MEJOR_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def _autoajustar_columnas(ws) -> None:
    for col_cells in ws.columns:
        longitud = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(longitud + 3, 45)


def _escribir_encabezado(ws, fila: int, valores: list[str]) -> None:
    for col, valor in enumerate(valores, start=1):
        celda = ws.cell(row=fila, column=col, value=valor)
        celda.font = ENCABEZADO_FONT
        celda.fill = ENCABEZADO_FILL
        celda.alignment = Alignment(horizontal="center")


class ExcelReportGenerator(ReportGenerator):
    extension = "xlsx"

    def generar(self, contexto: ReportContext, output_path: Path) -> Path:
        wb = Workbook()
        lang = contexto.lang

        self._hoja_resumen(wb.active, contexto)
        self._hoja_cv_por_fold(wb.create_sheet(get_text("reports.cv_sheet", lang)), contexto)
        self._hoja_hiperparametros(wb.create_sheet(get_text("reports.hyperparams_sheet", lang)), contexto)
        if contexto.pruebas_estadisticas:
            self._hoja_pruebas_estadisticas(wb.create_sheet(get_text("reports.stats_sheet", lang)), contexto)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path

    def _hoja_resumen(self, ws, contexto: ReportContext) -> None:
        lang = contexto.lang
        ws.title = get_text("reports.section2", lang)
        ws["A1"] = f"{get_text('reports.cover_title', lang)} - {get_text('reports.cover_subtitle', lang)}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"{get_text('reports.cover_date', lang)}{contexto.fecha_generacion}"
        ws["A3"] = f"{get_text('reports.executive_best', lang)}: {contexto.resumen_entrenamiento.get('mejor_modelo')} (F1-macro {get_text('reports.metric_mean', lang)}: {contexto.resumen_entrenamiento.get('f1_macro_cv_promedio', ''):.4f})"
        ws["A3"].font = Font(bold=True)

        tabla = contexto.tabla_comparativa
        metricas = [c for c in tabla.columns if c[1] == "mean"]
        stds = [c for c in tabla.columns if c[1] == "std"]

        fila_inicio = 5
        _escribir_encabezado(ws, fila_inicio,
                             [get_text("reports.models_header", lang)]
                             + [f"{m[0]} {get_text('reports.metric_mean', lang)}" for m in metricas]
                             + [f"{s[0]} {get_text('reports.metric_std', lang)}" for s in stds])
        mejor_modelo = contexto.resumen_entrenamiento.get("mejor_modelo")
        for i, modelo in enumerate(tabla.index, start=fila_inicio + 1):
            ws.cell(row=i, column=1, value=modelo)
            for j, m in enumerate(metricas, start=2):
                ws.cell(row=i, column=j, value=round(float(tabla.loc[modelo, m]), 4))
            for j, s in enumerate(stds, start=2 + len(metricas)):
                ws.cell(row=i, column=j, value=round(float(tabla.loc[modelo, s]), 4))
            if modelo == mejor_modelo:
                for col in range(1, 2 + len(metricas) + len(stds)):
                    ws.cell(row=i, column=col).fill = MEJOR_FILL
        _autoajustar_columnas(ws)

    def _hoja_cv_por_fold(self, ws, contexto: ReportContext) -> None:
        for r_idx, row in enumerate(dataframe_to_rows(contexto.cv_resultados_por_fold, index=False, header=True), start=1):
            for c_idx, valor in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=valor)
        _escribir_encabezado(ws, 1, list(contexto.cv_resultados_por_fold.columns))
        _autoajustar_columnas(ws)

    def _hoja_hiperparametros(self, ws, contexto: ReportContext) -> None:
        lang = contexto.lang
        _escribir_encabezado(ws, 1,
                             [get_text("reports.models_header", lang),
                              get_text("reports.hyperparam_header", lang),
                              get_text("reports.hyperparam_value", lang)])
        fila = 2
        for modelo, info in contexto.hiperparametros.items():
            params = info.get("params", {})
            for k, v in params.items():
                ws.cell(row=fila, column=1, value=modelo)
                ws.cell(row=fila, column=2, value=str(k))
                ws.cell(row=fila, column=3, value=str(v))
                fila += 1
        _autoajustar_columnas(ws)

    def _hoja_pruebas_estadisticas(self, ws, contexto: ReportContext) -> None:
        lang = contexto.lang
        pruebas = contexto.pruebas_estadisticas
        ws["A1"] = get_text("reports.friedman_text", lang)
        ws["A1"].font = Font(bold=True)
        friedman = pruebas["friedman"]
        ws["A2"] = get_text("reports.stats_chi2", lang)
        ws["B2"] = round(friedman["estadistico_chi2"], 4)
        ws["A3"] = get_text("reports.pvalue", lang)
        ws["B3"] = round(friedman["p_value"], 4)
        ws["A4"] = get_text("reports.stats_sig", lang)
        ws["B4"] = (get_text("reports.significant_yes", lang)
                    if friedman["significativo_alpha_0.05"]
                    else get_text("reports.significant_no", lang))

        fila = 6
        ws.cell(row=fila, column=1, value=get_text("reports.wilcoxon_title", lang)).font = Font(bold=True)
        fila += 1
        _escribir_encabezado(ws, fila,
                             [get_text("reports.wilcoxon_header_a", lang),
                              get_text("reports.wilcoxon_header_b", lang),
                              get_text("reports.wilcoxon_p", lang),
                              get_text("reports.wilcoxon_holm", lang),
                              get_text("reports.wilcoxon_sig", lang)])
        for r in pruebas["wilcoxon_pareado_holm"]:
            fila += 1
            ws.cell(row=fila, column=1, value=r["modelo_a"])
            ws.cell(row=fila, column=2, value=r["modelo_b"])
            ws.cell(row=fila, column=3, value=round(r["p_value"], 4))
            ws.cell(row=fila, column=4, value=round(r["p_value_holm"], 4))
            ws.cell(row=fila, column=5,
                    value=(get_text("reports.significant_yes", lang)
                           if r["significativo_holm_0.05"]
                           else get_text("reports.significant_no", lang)))
        _autoajustar_columnas(ws)
