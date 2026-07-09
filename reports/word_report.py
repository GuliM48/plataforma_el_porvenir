from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from dashboard.i18n import get_text
from reports.base import ReportGenerator
from reports.context import ReportContext


class WordReportGenerator(ReportGenerator):
    extension = "docx"

    def generar(self, contexto: ReportContext, output_path: Path) -> Path:
        doc = Document()
        self._estilos_base(doc)

        self._portada(doc, contexto)
        doc.add_page_break()
        self._resumen_ejecutivo(doc, contexto)
        self._tabla_comparativa(doc, contexto)
        self._figuras(doc, contexto)
        self._hiperparametros(doc, contexto)
        if contexto.pruebas_estadisticas:
            self._pruebas_estadisticas(doc, contexto)
        self._conclusiones(doc, contexto)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return output_path

    def _estilos_base(self, doc: Document) -> None:
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)

    def _portada(self, doc: Document, contexto: ReportContext) -> None:
        lang = contexto.lang
        titulo = doc.add_heading(
            f"{get_text('reports.cover_title', lang)}\n{get_text('reports.cover_subtitle', lang)}", level=0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub = doc.add_paragraph(get_text("reports.cover_desc", lang))
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.runs[0].font.size = Pt(14)
        sub.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        fecha = doc.add_paragraph(
            f"{get_text('reports.cover_date', lang)}{contexto.fecha_generacion}")
        fecha.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _resumen_ejecutivo(self, doc: Document, contexto: ReportContext) -> None:
        lang = contexto.lang
        doc.add_heading(get_text("reports.section1", lang), level=1)
        mejor = contexto.resumen_entrenamiento.get("mejor_modelo")
        f1 = contexto.resumen_entrenamiento.get("f1_macro_cv_promedio")
        cfg = contexto.resumen_entrenamiento.get("config", {})
        doc.add_paragraph(
            f"{get_text('reports.executive_summary', lang)} "
            f"(n_splits={cfg.get('n_splits')}, n_repeats={cfg.get('n_repeats')})."
        )
        p = doc.add_paragraph()
        p.add_run(f"{get_text('reports.executive_best', lang)} ").bold = True
        p.add_run(f"{mejor} (F1-macro = {f1:.4f}).").bold = True

    def _tabla_comparativa(self, doc: Document, contexto: ReportContext) -> None:
        lang = contexto.lang
        doc.add_heading(get_text("reports.section2", lang), level=1)
        tabla_df = contexto.tabla_comparativa
        metricas = [c for c in tabla_df.columns if c[1] == "mean"]

        tabla = doc.add_table(rows=1, cols=len(metricas) + 1)
        tabla.style = "Light Grid Accent 1"
        hdr = tabla.rows[0].cells
        hdr[0].text = get_text("reports.models_header", lang)
        for i, m in enumerate(metricas, start=1):
            hdr[i].text = m[0]

        for modelo in tabla_df.index:
            fila = tabla.add_row().cells
            fila[0].text = str(modelo)
            for i, m in enumerate(metricas, start=1):
                fila[i].text = f"{tabla_df.loc[modelo, m]:.4f}"
        doc.add_paragraph()

    def _figuras(self, doc: Document, contexto: ReportContext) -> None:
        lang = contexto.lang
        doc.add_heading(get_text("reports.section3", lang), level=1)
        titulos = {
            "roc": get_text("reports.fig_roc", lang),
            "matrices_confusion": get_text("reports.fig_confusion", lang),
            "importancia_variables": get_text("reports.fig_importance", lang),
            "comparacion_modelos": get_text("reports.fig_comparison", lang),
            "boxplot_estadistico": get_text("reports.fig_boxplot", lang),
        }
        for clave, titulo in titulos.items():
            ruta = contexto.figuras.get(clave)
            if ruta is None:
                continue
            doc.add_heading(titulo, level=2)
            doc.add_picture(str(ruta), width=Inches(6.3))

    def _hiperparametros(self, doc: Document, contexto: ReportContext) -> None:
        lang = contexto.lang
        doc.add_heading(get_text("reports.section4", lang), level=1)
        for modelo, info in contexto.hiperparametros.items():
            if modelo == "ga_mlp":
                continue
            doc.add_paragraph(f"{modelo}: {info.get('params')}")

    def _pruebas_estadisticas(self, doc: Document, contexto: ReportContext) -> None:
        lang = contexto.lang
        doc.add_heading(get_text("reports.section5", lang), level=1)
        friedman = contexto.pruebas_estadisticas["friedman"]
        sig = (get_text("reports.significant_yes", lang)
               if friedman["significativo_alpha_0.05"]
               else get_text("reports.significant_no", lang))
        doc.add_paragraph(
            f"{get_text('reports.friedman_text', lang)}: "
            f"χ²={friedman['estadistico_chi2']:.4f}, "
            f"p={friedman['p_value']:.4f} "
            f"({sig} "
            f"a α=0.05, con {friedman['n_folds']} folds)."
        )
        if "advertencia" in friedman:
            p = doc.add_paragraph(friedman["advertencia"])
            p.runs[0].italic = True

        doc.add_paragraph(get_text("reports.wilcoxon_title", lang))
        for r in contexto.pruebas_estadisticas["wilcoxon_pareado_holm"]:
            marca = (f" ({get_text('reports.significant_yes', lang)})"
                     if r["significativo_holm_0.05"] else "")
            doc.add_paragraph(
                f"  \u2022 {r['modelo_a']} vs {r['modelo_b']}: "
                f"p_holm={r['p_value_holm']:.4f}{marca}",
                style="List Bullet")

    def _conclusiones(self, doc: Document, contexto: ReportContext) -> None:
        lang = contexto.lang
        doc.add_heading(get_text("reports.section6", lang), level=1)
        mejor = contexto.resumen_entrenamiento.get("mejor_modelo")
        doc.add_paragraph(
            f"{get_text('reports.executive_best', lang)} {mejor} "
            f"{get_text('reports.conclusions', lang)}"
        )
